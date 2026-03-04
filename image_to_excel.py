"""
Image Text Extraction to Excel
특정 폴더의 이미지 파일들을 읽어서 OCR로 텍스트를 추출하고 Excel에 저장하는 프로그램
"""

import os
import sys
import argparse
from datetime import datetime
from pathlib import Path

try:
    from PIL import Image
    import pytesseract
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError as e:
    print(f"필수 라이브러리가 없습니다. 다음 명령어로 설치하세요:")
    print("pip install Pillow pytesseract openpyxl")
    sys.exit(1)

# 지원 이미지 확장자
SUPPORTED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif", ".gif", ".webp"}


def get_image_files(folder_path: str) -> list[Path]:
    """폴더에서 이미지 파일 목록을 반환합니다."""
    folder = Path(folder_path)
    if not folder.exists():
        raise FileNotFoundError(f"폴더를 찾을 수 없습니다: {folder_path}")
    if not folder.is_dir():
        raise NotADirectoryError(f"폴더 경로가 아닙니다: {folder_path}")

    image_files = sorted(
        [f for f in folder.iterdir() if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS]
    )
    return image_files


def extract_text_from_image(image_path: Path, lang: str = "kor+eng") -> tuple[str, str]:
    """
    이미지에서 텍스트를 추출합니다.
    반환값: (추출된 텍스트, 오류 메시지)
    """
    try:
        image = Image.open(image_path)
        # 이미지가 RGBA나 P 모드인 경우 RGB로 변환
        if image.mode in ("RGBA", "P"):
            image = image.convert("RGB")

        # OCR 설정: PSM 3 (자동 페이지 세그멘테이션, OSD 없음)
        custom_config = r"--oem 3 --psm 3"
        text = pytesseract.image_to_string(image, lang=lang, config=custom_config)
        return text.strip(), ""
    except pytesseract.TesseractNotFoundError:
        return "", "Tesseract가 설치되지 않았습니다. https://github.com/tesseract-ocr/tesseract 참조"
    except Exception as e:
        return "", str(e)


def apply_header_style(cell):
    """헤더 셀에 스타일을 적용합니다."""
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="FFFFFF")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def apply_data_style(cell, row_idx: int, wrap: bool = False):
    """데이터 셀에 스타일을 적용합니다."""
    bg_color = "EBF3FB" if row_idx % 2 == 0 else "FFFFFF"
    cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    thin = Side(style="thin", color="CCCCCC")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def save_to_excel(results: list[dict], output_path: str):
    """추출된 텍스트 결과를 Excel 파일로 저장합니다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "OCR 결과"

    # 헤더 정의
    headers = ["#", "파일명", "파일 경로", "추출된 텍스트", "문자 수", "처리 시각", "오류"]
    ws.append(headers)

    for cell in ws[1]:
        apply_header_style(cell)
    ws.row_dimensions[1].height = 30

    # 데이터 행 추가
    for i, result in enumerate(results, start=1):
        row = [
            i,
            result["filename"],
            result["filepath"],
            result["text"],
            len(result["text"]),
            result["timestamp"],
            result["error"],
        ]
        ws.append(row)

        row_idx = i + 1  # 헤더가 1행이므로 데이터는 2행부터
        for col_idx, cell in enumerate(ws[row_idx], start=1):
            is_text_col = (col_idx == 4)  # "추출된 텍스트" 열
            apply_data_style(cell, i, wrap=is_text_col)

        # 텍스트가 많은 경우 행 높이 조정
        text_lines = result["text"].count("\n") + 1
        ws.row_dimensions[row_idx].height = min(max(text_lines * 15, 25), 300)

    # 열 너비 설정
    column_widths = {"A": 5, "B": 30, "C": 45, "D": 60, "E": 10, "F": 22, "G": 30}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 요약 시트 추가
    ws_summary = wb.create_sheet(title="요약")
    summary_data = [
        ["항목", "값"],
        ["총 이미지 수", len(results)],
        ["텍스트 추출 성공", sum(1 for r in results if r["text"])],
        ["텍스트 추출 실패 / 빈 결과", sum(1 for r in results if not r["text"])],
        ["오류 발생", sum(1 for r in results if r["error"])],
        ["처리 시각", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ]
    for row in summary_data:
        ws_summary.append(row)

    for cell in ws_summary[1]:
        apply_header_style(cell)
    ws_summary.column_dimensions["A"].width = 25
    ws_summary.column_dimensions["B"].width = 30

    wb.save(output_path)
    print(f"\nExcel 파일 저장 완료: {output_path}")


def process_folder(folder_path: str, output_path: str = None, lang: str = "kor+eng"):
    """
    폴더 내 모든 이미지를 처리하여 Excel로 저장합니다.

    Args:
        folder_path: 이미지가 있는 폴더 경로
        output_path: 저장할 Excel 파일 경로 (None이면 자동 생성)
        lang: OCR 언어 (기본: 한국어+영어)
    """
    print(f"이미지 폴더: {folder_path}")
    print(f"OCR 언어: {lang}")
    print("-" * 50)

    # 이미지 파일 목록 가져오기
    image_files = get_image_files(folder_path)

    if not image_files:
        print("처리할 이미지 파일이 없습니다.")
        print(f"지원 형식: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
        return

    print(f"발견된 이미지: {len(image_files)}개\n")

    # 출력 파일 경로 설정
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join(folder_path, f"OCR_결과_{timestamp}.xlsx")

    # 각 이미지에서 텍스트 추출
    results = []
    for idx, image_path in enumerate(image_files, start=1):
        print(f"[{idx:3d}/{len(image_files)}] 처리 중: {image_path.name}")
        text, error = extract_text_from_image(image_path, lang=lang)

        if error:
            print(f"          오류: {error}")
        elif text:
            preview = text[:80].replace("\n", " ")
            print(f"          추출: {preview}{'...' if len(text) > 80 else ''}")
        else:
            print(f"          결과: 텍스트 없음")

        results.append({
            "filename": image_path.name,
            "filepath": str(image_path.resolve()),
            "text": text,
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "error": error,
        })

    # Excel로 저장
    save_to_excel(results, output_path)

    # 최종 통계
    success = sum(1 for r in results if r["text"])
    print(f"\n처리 완료!")
    print(f"  - 전체: {len(results)}개")
    print(f"  - 텍스트 추출 성공: {success}개")
    print(f"  - 텍스트 없음 / 오류: {len(results) - success}개")


def main():
    parser = argparse.ArgumentParser(
        description="이미지 파일에서 텍스트를 추출하여 Excel로 저장합니다.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
사용 예시:
  python image_to_excel.py ./images
  python image_to_excel.py ./images -o 결과.xlsx
  python image_to_excel.py ./images -o 결과.xlsx --lang eng
  python image_to_excel.py ./images --lang kor+eng+jpn
        """,
    )
    parser.add_argument(
        "folder",
        help="이미지 파일이 있는 폴더 경로",
    )
    parser.add_argument(
        "-o", "--output",
        default=None,
        help="저장할 Excel 파일 경로 (기본: 이미지 폴더 내 OCR_결과_[날짜시간].xlsx)",
    )
    parser.add_argument(
        "--lang",
        default="kor+eng",
        help="OCR 언어 코드 (기본: kor+eng). 예: eng, kor, kor+eng+jpn",
    )

    args = parser.parse_args()

    try:
        process_folder(
            folder_path=args.folder,
            output_path=args.output,
            lang=args.lang,
        )
    except (FileNotFoundError, NotADirectoryError) as e:
        print(f"오류: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
